# Copyright (c) 2025 longooc.com
# Author: Bin Lee
# Email: blee@longooc.com

"""Generate DOCX templates for CTD modules from the IND 2.4/2.6 template."""

from __future__ import annotations

import json
from collections import OrderedDict
from itertools import count
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

_DEFAULT_TEMPLATE_PATHS = [
    Path(__file__).resolve().parents[2] / "ncd" / "ind_24_26_template.json",
]

_MODULE_KEY_MAP = {
    "2.4": "Module 2.4 Nonclinical Overview",
    "2.6-written": "Module 2.6 Written Summary",
    "2.6-tabulated": "Module 2.6 Tabulated Summary",
}

_SDT_ID = count(1)


def _ensure_docx() -> Tuple[Any, Any]:
    try:
        from docx import Document
        from docx.shared import RGBColor
    except ImportError as exc:  # pragma: no cover - handled at runtime
        raise RuntimeError(
            "python-docx is required to generate DOCX templates. "
            "Install it with `poetry add python-docx`."
        ) from exc
    return Document, RGBColor


def load_module_entries(
    module_key: str,
    template_path: Path | None = None,
) -> Tuple[str, List[Dict[str, Any]]]:
    payload = _load_template_payload(template_path)
    resolved_key = _resolve_module_key(module_key, payload)
    raw_entries = payload.get(resolved_key)
    if not isinstance(raw_entries, list):
        raise ValueError(f"Template module '{resolved_key}' is missing or not a list.")
    entries = [entry for entry in raw_entries if isinstance(entry, dict)]
    if not entries:
        raise ValueError(f"No template entries found for module '{resolved_key}'.")
    return resolved_key, entries


def write_module_docx(
    module_key: str,
    output_path: Path,
    template_path: Path | None = None,
) -> None:
    Document, RGBColor = _ensure_docx()
    module_title, entries = load_module_entries(module_key, template_path)
    _write_docx(module_key, module_title, entries, output_path, Document, RGBColor)


def write_section_docx(
    section_number: str,
    output_path: Path,
    template_path: Path | None = None,
    module_key: str | None = None,
) -> None:
    Document, RGBColor = _ensure_docx()
    resolved_key, module_title, entries = _resolve_section_entries(
        section_number, template_path, module_key
    )
    _write_docx(resolved_key, module_title, entries, output_path, Document, RGBColor)


def _load_template_payload(template_path: Path | None) -> Dict[str, Any]:
    paths = [template_path] if template_path else _DEFAULT_TEMPLATE_PATHS
    for path in paths:
        if not path:
            continue
        if not path.exists():
            continue
        return json.loads(path.read_text(encoding="utf-8"))
    raise FileNotFoundError("Unable to find ind_24_26_template.json.")


def _resolve_module_key(module_key: str, payload: Dict[str, Any]) -> str:
    cleaned = module_key.strip()
    if cleaned in payload:
        return cleaned
    mapped = _MODULE_KEY_MAP.get(cleaned.lower(), cleaned)
    if mapped in payload:
        return mapped
    raise KeyError(f"Unknown module key '{module_key}'.")


def _resolve_section_entries(
    section_number: str,
    template_path: Path | None,
    module_key: str | None = None,
) -> Tuple[str, str, List[Dict[str, Any]]]:
    section_number = str(section_number).strip()
    if not section_number:
        raise ValueError("Section number is required.")
    if module_key:
        module_title, entries = load_module_entries(module_key, template_path)
        selected = _filter_entries_by_section(entries, section_number)
        if not selected:
            raise ValueError(
                f"No template entries found for section '{section_number}' "
                f"in module '{module_key}'."
            )
        return module_key, module_title, selected

    for candidate in ("2.6-written", "2.6-tabulated", "2.4"):
        try:
            module_title, entries = load_module_entries(candidate, template_path)
        except (KeyError, ValueError):
            continue
        selected = _filter_entries_by_section(entries, section_number)
        if selected:
            return candidate, module_title, selected

    raise ValueError(f"No template entries found for section '{section_number}'.")


def _group_entries_by_section(
    entries: Iterable[Dict[str, Any]],
) -> "OrderedDict[str, Dict[str, Any]]":
    grouped: "OrderedDict[str, Dict[str, Any]]" = OrderedDict()
    for entry in entries:
        section_number = str(entry.get("Section") or "").strip()
        if not section_number:
            continue
        section_title = str(entry.get("Section Header") or "").strip()
        if section_number not in grouped:
            grouped[section_number] = {"title": section_title, "entries": []}
        grouped[section_number]["entries"].append(entry)
    return grouped


def _filter_entries_by_section(
    entries: Iterable[Dict[str, Any]],
    section_number: str,
) -> List[Dict[str, Any]]:
    return [
        entry
        for entry in entries
        if str(entry.get("Section") or "").strip() == section_number
    ]


def _write_docx(
    module_key: str,
    module_title: str,
    entries: Iterable[Dict[str, Any]],
    output_path: Path,
    Document: Any,
    RGBColor: Any,
) -> None:
    sections = _group_entries_by_section(entries)

    doc = Document()
    _set_page_header(doc, module_key, module_title)
    doc.add_heading(module_title, level=0)

    for section_number, section in sections.items():
        heading = _format_section_heading(section_number, section["title"])
        doc.add_heading(heading, level=1)
        for entry in section["entries"]:
            element_heading = _format_element_heading(entry)
            if element_heading:
                doc.add_heading(element_heading, level=2)
            content = str(entry.get("Content") or "").strip()
            if content:
                _add_content_control(doc, content, RGBColor)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def _format_section_heading(section_number: str, section_title: str) -> str:
    if section_title:
        return f"{section_number} {section_title}"
    return section_number


def _format_element_heading(entry: Dict[str, Any]) -> str:
    element = str(entry.get("Subsection Element Numbering") or "").strip()
    if element:
        return element
    subsection = str(entry.get("Subsection") or "").strip()
    if subsection:
        return subsection
    return ""


def _set_page_header(doc: Any, module_key: str, module_title: str) -> None:
    header_text = _format_page_header(module_key, module_title)
    if not header_text:
        return
    for section in doc.sections:
        header = section.header
        if header.paragraphs:
            paragraph = header.paragraphs[0]
            paragraph.text = header_text
        else:
            header.add_paragraph(header_text)


def _format_page_header(module_key: str, module_title: str) -> str:
    key = module_key.strip().lower()
    title = module_title.strip()
    if key.startswith("2.4") or "2.4" in title:
        return "2.4 Nonclinical Overview"
    if key.startswith("2.6") or "2.6" in title:
        if "Tabulated" in title:
            return "2.6 Tabulated Summary"
        if "Written" in title:
            return "2.6 Written Summary"
        return "2.6 Nonclinical Summary"
    if title.lower().startswith("module "):
        return title[7:]
    return title


def _apply_placeholder_style(run: Any, doc: Any, rgb_color: Any) -> None:
    try:
        run.style = doc.styles["Placeholder Text"]
        return
    except Exception:
        pass
    run.font.color.rgb = rgb_color(0x80, 0x80, 0x80)
    run.font.italic = True


def _add_content_control(doc: Any, text: str, rgb_color: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    _apply_placeholder_style(run, doc, rgb_color)

    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")
    sdt.append(sdt_pr)

    sdt_id = OxmlElement("w:id")
    sdt_id.set(qn("w:val"), str(next(_SDT_ID)))
    sdt_pr.append(sdt_id)

    sdt_text = OxmlElement("w:text")
    sdt_pr.append(sdt_text)

    showing_placeholder = OxmlElement("w:showingPlcHdr")
    sdt_pr.append(showing_placeholder)

    sdt_content = OxmlElement("w:sdtContent")
    sdt.append(sdt_content)

    p = paragraph._p
    parent = p.getparent()
    if parent is None:
        return
    index = parent.index(p)
    parent.remove(p)
    sdt_content.append(p)
    parent.insert(index, sdt)
