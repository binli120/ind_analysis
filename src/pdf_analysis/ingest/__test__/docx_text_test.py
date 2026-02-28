# Copyright (c) 2025 longooc.com
# Author: Bin Lee
# Email: blee@longooc.com

"""Unit tests for pdf_analysis.ingest.docx_text helpers."""

from __future__ import annotations

import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[4]
SRC_PATH = REPO_ROOT / "src"
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
if str(SRC_PATH) not in sys.path:
    sys.path.insert(0, str(SRC_PATH))

from pdf_analysis.ingest import docx_text


def _write_docx(path: Path) -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("First paragraph")
    doc.add_paragraph("Second paragraph")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    doc.save(path)


def test_extract_docx_pages_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _write_docx(path)

    pages = docx_text.extract_docx_pages(path, max_chars=50)
    assert pages
    assert pages[0]["page_number"] == 1
    assert "First paragraph" in pages[0]["text"]

    tables = docx_text.extract_docx_tables(path)
    assert len(tables) == 1
    df = tables[0]["dataframe"]
    assert list(df.columns) == ["Header A", "Header B"]
    assert df.iloc[0].tolist() == ["Value 1", "Value 2"]
