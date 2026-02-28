# Copyright (c) 2025 longooc.com
# Author: Bin Lee
# Email: blee@longooc.com

"""DOCX text + table extraction helpers."""

# @author: Bin Lee
# @email: blee@longooc.com

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

import pandas as pd

try:
    from docx import Document
except ModuleNotFoundError:  # pragma: no cover - optional dependency
    Document = None  # type: ignore[assignment]

_DOCX_PAGE_MAX_CHARS = int(os.getenv("DOCX_PAGE_MAX_CHARS", "2000"))
_WHITESPACE_RE = re.compile(r"\s+")


def _load_docx(docx_path: Path):
    if Document is None:  # pragma: no cover - optional dependency
        raise ModuleNotFoundError(
            "python-docx is required for DOCX extraction. "
            "Install it with `poetry add python-docx`."
        )
    return Document(str(docx_path))


def _normalize_cell(text: str) -> str:
    cleaned = _WHITESPACE_RE.sub(" ", text or "").strip()
    return cleaned


def _dedupe_headers(headers: Sequence[str]) -> List[str]:
    counts: dict[str, int] = {}
    unique: List[str] = []
    for header in headers:
        base = header or "column"
        count = counts.get(base, 0) + 1
        counts[base] = count
        if count == 1:
            unique.append(base)
        else:
            unique.append(f"{base}_{count}")
    return unique


def _iter_table_rows(table) -> Iterable[List[str]]:
    for row in table.rows:
        cells = [_normalize_cell(cell.text) for cell in row.cells]
        if any(cells):
            yield cells


def extract_docx_tables(docx_path: Path) -> List[Dict[str, Any]]:
    doc = _load_docx(docx_path)
    tables: List[Dict[str, Any]] = []
    for idx, table in enumerate(doc.tables, start=1):
        rows = list(_iter_table_rows(table))
        if not rows:
            continue
        max_cols = max(len(row) for row in rows)
        padded = [row + [""] * (max_cols - len(row)) for row in rows]
        header = padded[0]
        data_rows = padded[1:] if len(padded) > 1 else []
        if not any(header):
            header = [f"Column {col_idx + 1}" for col_idx in range(max_cols)]
            data_rows = padded
        header = _dedupe_headers([_normalize_cell(h) for h in header])
        df = pd.DataFrame(data_rows, columns=header)
        tables.append(
            {
                "page_number": 1,
                "index_on_page": idx,
                "engine": "docx",
                "dataframe": df,
                "bbox": None,
            }
        )
    return tables


def _iter_table_text_lines(tables) -> Iterable[str]:
    for table in tables:
        for row in _iter_table_rows(table):
            yield " | ".join(row)


def extract_docx_pages(
    docx_path: Path,
    *,
    max_chars: int = _DOCX_PAGE_MAX_CHARS,
    max_pages: Optional[int] = None,
) -> List[Dict[str, Any]]:
    doc = _load_docx(docx_path)
    paragraphs: List[str] = []
    for paragraph in doc.paragraphs:
        text = _normalize_cell(paragraph.text)
        if text:
            paragraphs.append(text)
    if not paragraphs:
        paragraphs = list(_iter_table_text_lines(doc.tables))

    if not paragraphs:
        return [{"page_number": 1, "text": ""}]

    pages: List[Dict[str, Any]] = []
    buffer: List[str] = []
    current_len = 0
    page_number = 1
    for paragraph in paragraphs:
        if buffer and current_len + len(paragraph) + 2 > max_chars:
            pages.append({"page_number": page_number, "text": "\n\n".join(buffer)})
            page_number += 1
            if max_pages and page_number > max_pages:
                return pages
            buffer = [paragraph]
            current_len = len(paragraph)
            continue
        buffer.append(paragraph)
        current_len += len(paragraph) + 2

    if buffer and (not max_pages or len(pages) < max_pages):
        pages.append({"page_number": page_number, "text": "\n\n".join(buffer)})
    return pages
